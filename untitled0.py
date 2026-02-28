import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog
from PIL import Image, ImageTk
import os, sys
import vlc
import PyPDF2
import pyttsx3
import docx

class FileOpener:
    def __init__(self, root):
        self.root = root
        self.root.title("Multi File Opener")
        self.root.geometry("1000x700")

        # VLC setup
        self.vlc_instance = vlc.Instance()
        self.vlc_player = self.vlc_instance.media_player_new()
        self.video_loaded = False

        # Text-to-Speech setup
        self.engine = pyttsx3.init()
        self.is_reading = False  # track if speaking

        # ---------------- Text Frame ----------------
        self.text_frame = tk.Frame(root)
        self.text_area = tk.Text(self.text_frame, wrap="word")
        self.scrollbar = tk.Scrollbar(self.text_frame, command=self.text_area.yview)
        self.text_area.config(yscrollcommand=self.scrollbar.set)
        self.text_area.pack(side=tk.LEFT, expand=True, fill="both")
        self.scrollbar.pack(side=tk.RIGHT, fill="y")

        self.text_controls = tk.Frame(self.text_frame, bg="lightgray")
        self.btn_audio_start = tk.Button(self.text_controls, text="🔊 Read Aloud", command=self.read_aloud)
        self.btn_audio_stop = tk.Button(self.text_controls, text="⏹ Stop Reading", command=self.stop_reading)
        self.btn_audio_start.pack(side=tk.LEFT, padx=5)
        self.btn_audio_stop.pack(side=tk.LEFT, padx=5)
        self.text_controls.pack(fill="x")

        # ---------------- Video Frame ----------------
        self.video_frame = tk.Frame(root, bg="black")
        self.video_canvas = tk.Canvas(self.video_frame, bg="black")
        self.video_canvas.pack(expand=True, fill="both")

        # Video Controls
        self.controls = tk.Frame(self.video_frame, bg="gray")
        self.btn_play = tk.Button(self.controls, text="▶ Play", command=self.play_video)
        self.btn_pause = tk.Button(self.controls, text="⏸ Pause", command=self.pause_video)
        self.btn_stop = tk.Button(self.controls, text="⏹ Stop", command=self.stop_video)
        self.btn_play.pack(side=tk.LEFT, padx=5)
        self.btn_pause.pack(side=tk.LEFT, padx=5)
        self.btn_stop.pack(side=tk.LEFT, padx=5)

        self.timeline = tk.Scale(
            self.controls, from_=0, to=100, orient="horizontal", length=500, command=self.seek_video
        )
        self.timeline.pack(side=tk.LEFT, padx=10)

        self.time_label = tk.Label(self.controls, text="00:00 / 00:00", bg="gray", fg="white")
        self.time_label.pack(side=tk.LEFT, padx=5)

        self.controls.pack(fill="x")

        # ---------------- Image Frame ----------------
        self.image_frame = tk.Frame(root, bg="white")
        self.image_canvas = tk.Canvas(self.image_frame, bg="white")
        self.image_canvas.pack(expand=True, fill="both")
        self.zoom_factor = 1.0
        self.original_img = None

        self.img_controls = tk.Frame(self.image_frame, bg="lightgray")
        self.btn_zoom_in = tk.Button(self.img_controls, text="🔍 Zoom In", command=self.zoom_in)
        self.btn_zoom_out = tk.Button(self.img_controls, text="🔎 Zoom Out", command=self.zoom_out)
        self.btn_zoom_in.pack(side=tk.LEFT, padx=5)
        self.btn_zoom_out.pack(side=tk.LEFT, padx=5)
        self.img_controls.pack(fill="x")

        # ---------------- Menu ----------------
        menu_bar = tk.Menu(root)
        file_menu = tk.Menu(menu_bar, tearoff=0)
        file_menu.add_command(label="Open File", command=self.open_file)
        file_menu.add_command(label="Save Text", command=self.save_text)
        file_menu.add_separator()
        file_menu.add_command(label="Exit", command=root.quit)
        menu_bar.add_cascade(label="File", menu=file_menu)

        search_menu = tk.Menu(menu_bar, tearoff=0)
        search_menu.add_command(label="Search in Text", command=self.search_text)
        menu_bar.add_cascade(label="Search", menu=search_menu)

        root.config(menu=menu_bar)

        # Update timeline
        self.update_timeline()

    # ---------------- File Open ----------------
    def open_file(self):
        path = filedialog.askopenfilename(
            title="Select a file",
            filetypes=[("All Files", "*.*"),
                       ("Text Files", "*.txt"),
                       ("PDF Files", "*.pdf"),
                       ("Word Files", "*.docx"),
                       ("Video Files", "*.mp4;*.avi;*.mov;*.mkv"),
                       ("Image Files", "*.jpg;*.jpeg;*.png;*.gif;*.bmp")]
        )
        if not path:
            return

        ext = os.path.splitext(path)[1].lower()
        if ext == ".txt":
            self.show_text(path)
        elif ext == ".pdf":
            self.show_pdf(path)
        elif ext == ".docx":
            self.show_docx(path)
        elif ext in [".mp4", ".avi", ".mov", ".mkv"]:
            self.show_video(path)
        elif ext in [".jpg", ".jpeg", ".png", ".gif", ".bmp"]:
            self.show_image(path)
        else:
            messagebox.showerror("Error", f"Unsupported file: {ext}")

    # ---------------- Text / PDF / DOCX ----------------
    def show_text(self, path):
        self.hide_all()
        self.text_frame.pack(fill="both", expand=True)
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                data = f.read()
            self.text_area.delete("1.0", tk.END)
            self.text_area.insert(tk.END, data)
            self.current_file = path
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def show_pdf(self, path):
        self.hide_all()
        self.text_frame.pack(fill="both", expand=True)
        try:
            content = ""
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    content += page.extract_text() or ""
            self.text_area.delete("1.0", tk.END)
            self.text_area.insert(tk.END, content if content else "[No text found]")
            self.current_file = None
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def show_docx(self, path):
        self.hide_all()
        self.text_frame.pack(fill="both", expand=True)
        try:
            doc = docx.Document(path)
            content = "\n".join([p.text for p in doc.paragraphs])
            self.text_area.delete("1.0", tk.END)
            self.text_area.insert(tk.END, content if content else "[Empty document]")
            self.current_file = None
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def save_text(self):
        if hasattr(self, "current_file") and self.current_file:
            try:
                with open(self.current_file, "w", encoding="utf-8") as f:
                    f.write(self.text_area.get("1.0", tk.END))
                messagebox.showinfo("Save", "File saved.")
            except Exception as e:
                messagebox.showerror("Error", str(e))
        else:
            messagebox.showwarning("Save", "This file cannot be saved.")

    def search_text(self):
        word = simpledialog.askstring("Search", "Enter word to find:")
        if not word:
            return
        self.text_area.tag_remove("highlight", "1.0", tk.END)
        start = "1.0"
        found = 0
        while True:
            start = self.text_area.search(word, start, stopindex=tk.END, nocase=True)
            if not start:
                break
            end = f"{start}+{len(word)}c"
            self.text_area.tag_add("highlight", start, end)
            start = end
            found += 1
        self.text_area.tag_config("highlight", background="yellow")
        messagebox.showinfo("Search", f"Found {found} matches" if found else "Not found")

    def read_aloud(self):
        text = self.text_area.get("1.0", tk.END).strip()
        if text:
            self.is_reading = True
            self.engine.say(text)
            self.engine.startLoop(False)  # run async
        else:
            messagebox.showinfo("Read Aloud", "No text to read.")

    def stop_reading(self):
        if self.is_reading:
            self.engine.endLoop()
            self.is_reading = False

    # ---------------- Video ----------------
    def show_video(self, path):
        self.hide_all()
        self.video_frame.pack(fill="both", expand=True)
        media = self.vlc_instance.media_new(path)
        self.vlc_player.set_media(media)
        if os.name == "nt":
            self.vlc_player.set_hwnd(self.video_canvas.winfo_id())
        elif sys.platform == "darwin":
            self.vlc_player.set_nsobject(self.video_canvas.winfo_id())
        else:
            self.vlc_player.set_xwindow(self.video_canvas.winfo_id())
        self.vlc_player.play()
        self.video_loaded = True

    def play_video(self):
        if self.video_loaded: self.vlc_player.play()

    def pause_video(self):
        if self.video_loaded: self.vlc_player.pause()

    def stop_video(self):
        if self.video_loaded:
            self.vlc_player.stop()
            self.timeline.set(0)
            self.time_label.config(text="00:00 / 00:00")

    def update_timeline(self):
        if self.video_loaded and self.vlc_player.is_playing():
            length = self.vlc_player.get_length() // 1000
            pos = self.vlc_player.get_time() // 1000
            if length > 0:
                percent = (pos / length) * 100
                self.timeline.set(percent)
                self.time_label.config(text=f"{self.format_time(pos)} / {self.format_time(length)}")
        self.root.after(500, self.update_timeline)

    def seek_video(self, value):
        if self.video_loaded:
            length = self.vlc_player.get_length()
            if length > 0:
                new_time = int((float(value) / 100) * length)
                self.vlc_player.set_time(new_time)

    def format_time(self, sec):
        m, s = divmod(int(sec), 60)
        return f"{m:02}:{s:02}"

    # ---------------- Image ----------------
    def show_image(self, path):
        self.hide_all()
        self.image_frame.pack(fill="both", expand=True)
        try:
            self.original_img = Image.open(path)
            self.zoom_factor = 1.0
            self.display_image()
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def display_image(self):
        if self.original_img:
            w, h = self.original_img.size
            new_size = (int(w * self.zoom_factor), int(h * self.zoom_factor))
            img_resized = self.original_img.resize(new_size, Image.LANCZOS)  # ✅ fix
            self.tk_img = ImageTk.PhotoImage(img_resized)
            self.image_canvas.delete("all")
            self.image_canvas.create_image(0, 0, anchor="nw", image=self.tk_img)
            self.image_canvas.config(scrollregion=self.image_canvas.bbox("all"))

    def zoom_in(self):
        self.zoom_factor *= 1.2
        self.display_image()

    def zoom_out(self):
        self.zoom_factor /= 1.2
        self.display_image()

    # ---------------- Utils ----------------
    def hide_all(self):
        self.text_frame.pack_forget()
        self.video_frame.pack_forget()
        self.image_frame.pack_forget()

# Run App
root = tk.Tk()
app = FileOpener(root)
root.mainloop()
